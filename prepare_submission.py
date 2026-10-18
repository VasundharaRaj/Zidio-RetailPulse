from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt
from pptx import Presentation
from pptx.util import Inches, Pt

root = Path.cwd()
out = root / "outputs"
sub = root / "submission"
sub.mkdir(exist_ok=True)

def csvfile(name):
    p = out / name
    return pd.read_csv(p) if p.exists() else pd.DataFrame()

def add_table(doc, df, limit=8):
    if df.empty:
        doc.add_paragraph("No data available.")
        return
    df = df.head(limit)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Shading Accent 1"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if pd.isna(val) else str(val)

def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

# Read actual project outputs
forecast = csvfile("demand_forecast.csv")
segments = csvfile("customer_segments.csv")
churn = csvfile("churn_predictions.csv")
importance = csvfile("churn_feature_importance.csv")
inventory = csvfile("inventory_analysis.csv")
slow = csvfile("slow_moving_products.csv")
metrics_path = out / "churn_metrics.txt"
metrics = metrics_path.read_text(encoding="utf-8", errors="replace") if metrics_path.exists() else "Metrics file not found."

# Create project report
doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = DPt(10)

doc.add_heading("RetailPulse", 0)
doc.add_paragraph("Retail Analytics, Forecasting and Customer Insights")
doc.add_paragraph("Zidio Development Internship Project")
doc.add_paragraph("Project Report | 14 October 2026")

doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "RetailPulse is a Python-based retail analytics project. It processes retail "
    "transactions and produces a baseline demand forecast, customer segments, "
    "customer churn-risk estimates, product movement summaries, and an interactive "
    "Streamlit dashboard."
)

doc.add_heading("2. Business Objectives", level=1)
add_bullets(doc, [
    "Clean and prepare transaction data for analysis.",
    "Forecast monthly sales using a baseline forecasting method.",
    "Segment customers using Recency, Frequency and Monetary value (RFM).",
    "Estimate potential customer churn using a Random Forest classifier.",
    "Identify product sales movement and slow-moving products.",
    "Present results in an interactive dashboard."
])

doc.add_heading("3. Dataset and Data Preparation", level=1)
doc.add_paragraph(
    "The source Online Retail dataset contains 541,909 rows and 8 columns, "
    "covering 1 December 2010 to 9 December 2011. Fields include invoice details, "
    "product descriptions, quantities, dates, unit prices, customer identifiers "
    "and countries. Monetary values are treated as pounds sterling (£)."
)
add_bullets(doc, [
    "Removed duplicate rows and cleaned product descriptions.",
    "Parsed invoice dates and prepared customer identifiers.",
    "Created Sales as Quantity multiplied by UnitPrice.",
    "Created an IsCancelled indicator from invoice identifiers.",
    "Saved the cleaned dataset to outputs/retail_cleaned.csv."
])

doc.add_heading("4. Demand Forecasting", level=1)
doc.add_paragraph(
    "Positive-quantity and positive-price transactions are aggregated by month. "
    "The next three months are forecast using the average of the last three "
    "historical monthly sales totals. This is a baseline, not a production-grade "
    "forecasting model."
)
add_table(doc, forecast)
p = out / "demand_forecast.png"
if p.exists():
    doc.add_picture(str(p), width=DInches(6))
    doc.add_paragraph("Figure 1. Monthly sales forecast.")

doc.add_heading("5. Customer Segmentation", level=1)
doc.add_paragraph(
    "RFM features describe how recently a customer purchased, how frequently they "
    "purchased, and the total monetary value of their recorded purchases. The "
    "segmentation output groups customers using the implemented rules."
)
doc.add_paragraph("Customer records in output: " + str(len(segments)))
add_table(doc, segments)
p = out / "customer_segments.png"
if p.exists():
    doc.add_picture(str(p), width=DInches(5.8))
    doc.add_paragraph("Figure 2. Customer segmentation.")

doc.add_heading("6. Customer Churn Prediction", level=1)
doc.add_paragraph(
    "A Random Forest classifier estimates churn risk from customer-level historical "
    "features. The target uses inactivity during the final 90-day observation window "
    "as a proxy for churn. Inactivity does not necessarily mean a customer has "
    "permanently left."
)
doc.add_paragraph("Model evaluation output:")
doc.add_paragraph(metrics)
doc.add_paragraph("Feature importance results:")
add_table(doc, importance, 10)
p = out / "churn_feature_importance.png"
if p.exists():
    doc.add_picture(str(p), width=DInches(5.8))
    doc.add_paragraph("Figure 3. Churn model feature importance.")

doc.add_heading("7. Inventory and Product Movement", level=1)
doc.add_paragraph(
    "The analysis summarises product sales movement and recent observed sale dates "
    "and identifies slow-moving products. The source dataset does not contain "
    "actual stock-on-hand, supplier lead times or purchase orders, so this is "
    "not a complete inventory optimisation model."
)
doc.add_paragraph("Products analysed: " + str(len(inventory)))
add_table(doc, inventory, 10)
doc.add_paragraph("Examples of slow-moving products:")
add_table(doc, slow, 10)
for name, caption in [
    ("top_products.png", "Figure 4. Top products by sales movement."),
    ("monthly_units_sold.png", "Figure 5. Monthly units sold.")
]:
    p = out / name
    if p.exists():
        doc.add_picture(str(p), width=DInches(5.8))
        doc.add_paragraph(caption)

doc.add_heading("8. Interactive Dashboard", level=1)
doc.add_paragraph(
    "The Streamlit dashboard presents project outputs and visualisations. "
    "To start it from the project root, run:"
)
doc.add_paragraph("python -m streamlit run dashboard.py")

doc.add_heading("9. Business Recommendations", level=1)
add_bullets(doc, [
    "Compare the baseline forecast with seasonal and statistical models before making purchasing decisions.",
    "Review potentially at-risk customers and test suitable retention campaigns.",
    "Check slow-moving products against actual stock levels, margins and supplier lead times.",
    "Monitor model precision, recall and forecast error as new data becomes available.",
    "Refresh the analyses when newer validated retail data is available."
])

doc.add_heading("10. Limitations and Future Work", level=1)
add_bullets(doc, [
    "The dataset is historical and covers 2010–2011.",
    "The forecast is a simple baseline and does not explicitly model seasonality.",
    "The churn definition is a proxy based on an inactivity window.",
    "Inventory analysis does not include actual stock quantities or supplier information.",
    "Future work should compare models, validate on later time periods and monitor performance."
])

doc.add_heading("11. Conclusion", level=1)
doc.add_paragraph(
    "RetailPulse demonstrates an end-to-end workflow from data cleaning to analytics, "
    "visualisation and dashboard delivery. It provides a useful foundation for "
    "further development and more rigorous validation."
)
doc.save(sub / "RetailPulse_Project_Report.docx")

# Create PowerPoint presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def slide_title(slide, title, subtitle=""):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.7))
    para = box.text_frame.paragraphs[0]
    para.text = title
    para.font.size = Pt(28)
    para.font.bold = True
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(12), Inches(0.4))
        box2.text_frame.text = subtitle
        box2.text_frame.paragraphs[0].font.size = Pt(14)

def slide_bullets(title, items, subtitle=""):
    slide = prs.slides.add_slide(blank)
    slide_title(slide, title, subtitle)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.8), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = item
        para.font.size = Pt(21)
        para.space_after = Pt(16)
    return slide

def slide_image(title, filename, subtitle):
    p = out / filename
    if not p.exists():
        return
    slide = prs.slides.add_slide(blank)
    slide_title(slide, title, subtitle)
    slide.shapes.add_picture(str(p), Inches(1.1), Inches(1.55), width=Inches(11.1))

slide_bullets("RetailPulse", [
    "Retail Analytics, Forecasting and Customer Insights",
    "Zidio Development Internship Project",
    "Python, pandas, scikit-learn and Streamlit",
    "Demand forecasting, customer segmentation, churn and product movement"
])

slide_bullets("Business Objectives", [
    "Clean and prepare transaction data",
    "Create a baseline monthly sales forecast",
    "Segment customers using RFM",
    "Estimate customer churn risk",
    "Summarise product sales movement",
    "Display outputs in an interactive dashboard"
])

slide_bullets("Dataset and Preparation", [
    "541,909 rows and 8 original columns",
    "Historical dates: December 2010 to December 2011",
    "Removed duplicates and prepared date and description fields",
    "Created Sales and IsCancelled fields",
    "Saved cleaned transactions to outputs/retail_cleaned.csv"
])

slide_bullets("Demand Forecasting", [
    "Monthly aggregation of positive sales transactions",
    "Three-month forecast horizon",
    "Baseline uses the mean of the last three historical monthly totals",
    "Forecast should be compared with stronger models before operational use"
])
slide_image("Demand Forecast", "demand_forecast.png", "Monthly sales forecast output")
slide_image("Customer Segmentation", "customer_segments.png", "RFM-based customer grouping")
slide_image("Churn Model", "churn_feature_importance.png", "Random Forest feature importance")
slide_image("Product Movement", "top_products.png", "Top products by sales movement")
slide_image("Monthly Units Sold", "monthly_units_sold.png", "Monthly product movement")

slide_bullets("Dashboard and Deliverables", [
    "Interactive dashboard file: dashboard.py",
    "Run with: python -m streamlit run dashboard.py",
    "CSV outputs, metrics and charts are stored in outputs/",
    "Report and presentation are stored in submission/"
])

slide_bullets("Limitations and Next Steps", [
    "Historical data may not represent current retail behaviour",
    "Forecast is a simple baseline",
    "Churn uses an inactivity-based proxy definition",
    "No actual stock-on-hand data is available",
    "Next: model comparison, time-based validation and updated data"
])

prs.save(sub / "RetailPulse_Presentation.pptx")

readme = (
    "# RetailPulse\n\n"
    "Python retail analytics project covering data cleaning, demand forecasting, "
    "RFM customer segmentation, churn-risk estimation, product movement analysis "
    "and an interactive Streamlit dashboard.\n\n"
    "## Run the dashboard\n"
    "From the project root, install dependencies with "
    "`python -m pip install -r submission/requirements.txt`, then run "
    "`python -m streamlit run dashboard.py`.\n\n"
    "## Important limitations\n"
    "The data is historical (2010–2011). Forecasting is a simple baseline, churn "
    "uses an inactivity-window proxy, and inventory analysis does not include "
    "actual stock-on-hand or supplier lead times.\n"
)
(sub / "README.md").write_text(readme, encoding="utf-8")

requirements = (
    "pandas\nopenpyxl\nmatplotlib\nscikit-learn\nstreamlit\nplotly\n"
    "python-docx\npython-pptx\n"
)
(sub / "requirements.txt").write_text(requirements, encoding="utf-8")

print("\nSUCCESS: Submission files created.")
for p in sorted(sub.iterdir()):
    print(f"{p.name} - {p.stat().st_size:,} bytes")
